#!/usr/bin/env python3
"""นำเข้าค่าใช้จ่ายจาก Account12.xlsx

แหล่งความจริงของรายการจ่าย = Excel Comment บนเซลล์คอลัมน์「จ่าย」(H)
จับคู่จำนวนกับสูตร =SUM(...) ในเซลล์ เพื่อกันคอมเมนต์เก่าค้าง
ส่วนที่จับคู่ไม่ได้ / ไม่มีคอมเมนต์ → อื่นๆ ให้รวมเท่ากับยอดจ่าย
"""

from __future__ import annotations

import argparse
import re
import unicodedata
import uuid
from datetime import date, datetime
from pathlib import Path

import openpyxl
import psycopg

CATEGORY_SHEET = "expcat-sheet"
CATEGORY_OTHER = "expcat-other"
CATEGORY_PORK = "expcat-pork"
CATEGORY_MEATBALL = "expcat-meatball"
CATEGORY_VEG = "expcat-veg"
CATEGORY_MAKRO = "expcat-makro"

BRANCHES = {"หลุมดิน", "เจ็ดเสมียน"}


def iso(d: date | datetime) -> str:
    if isinstance(d, datetime):
        d = d.date()
    return d.isoformat()


def slug(text: str) -> str:
    s = unicodedata.normalize("NFKC", text)
    s = re.sub(r"\s+", "-", s.strip())
    s = re.sub(r"[^0-9A-Za-z\u0E00-\u0E7F\-]+", "", s)
    return (s[:48] or "item").lower()


def category_for(name: str) -> str:
    n = unicodedata.normalize("NFKC", name).strip()
    if n == "อื่นๆ":
        return CATEGORY_OTHER
    if "หมู" in n:
        return CATEGORY_PORK
    if "ลูกชิ้น" in n:
        return CATEGORY_MEATBALL
    if "ผัก" in n:
        return CATEGORY_VEG
    if n in {"M", "m"} or "แม็คโคร" in n or "แมคโคร" in n:
        return CATEGORY_MAKRO
    return CATEGORY_SHEET


def parse_comment(text: str) -> list[tuple[str, float]]:
    text = str(text or "")
    text = re.sub(r"^[^:\n]+:\s*\n?", "", text).strip()
    items: list[tuple[str, float]] = []
    for line in text.splitlines():
        line = unicodedata.normalize("NFKC", line).strip()
        if not line:
            continue
        m = re.match(
            r"^(.*?)[\s]+([0-9]{1,3}(?:,[0-9]{3})*|[0-9]+)(?:\.[0-9]+)?\s*$",
            line,
        )
        if not m:
            continue
        name = m.group(1).strip()
        amt = float(m.group(2).replace(",", ""))
        if name and amt >= 0:
            items.append((name, amt))
    return items


def parse_formula_nums(val) -> list[float] | None:
    if not isinstance(val, str):
        return None
    m = re.match(r"^=SUM\((.+)\)$", val.strip(), re.I)
    if not m:
        return None
    return [float(x) for x in re.findall(r"\d+(?:\.\d+)?", m.group(1))]


def merge_other(lines: list[tuple[str, float]]) -> list[tuple[str, float]]:
    named: list[tuple[str, float]] = []
    other = 0.0
    for n, a in lines:
        if n == "อื่นๆ":
            other += a
        else:
            named.append((n, a))
    if other >= 0.5:
        named.append(("อื่นๆ", round(other, 2)))
    return named


def match_lines(
    pay: float,
    formula_nums: list[float] | None,
    comment_items: list[tuple[str, float]],
) -> list[tuple[str, float]]:
    pay = round(pay, 2)
    if formula_nums and abs(sum(formula_nums) - pay) < 1:
        remaining = list(comment_items)
        lines: list[tuple[str, float]] = []
        for num in formula_nums:
            found = None
            for i, (n, a) in enumerate(remaining):
                if abs(a - num) < 0.01:
                    found = i
                    break
            if found is not None:
                lines.append(remaining.pop(found))
            else:
                lines.append(("อื่นๆ", num))
        return merge_other(lines)

    if comment_items:
        named = [(n, a) for n, a in comment_items if n != "อื่นๆ"]
        named_sum = round(sum(a for _, a in named), 2)
        other = round(pay - named_sum, 2)
        lines = list(named)
        if other >= 0.5:
            lines.append(("อื่นๆ", other))
        elif other < -0.5:
            return merge_other(comment_items)
        return merge_other(lines)

    return [("อื่นๆ", pay)] if pay > 0 else []


def parse_workbook(path: Path):
    wb_val = openpyxl.load_workbook(path, data_only=True)
    wb_form = openpyxl.load_workbook(path, data_only=False)
    days: dict[str, dict] = {}

    for sheet in wb_val.sheetnames:
        ws_v = wb_val[sheet]
        ws_f = wb_form[sheet]
        current: date | None = None
        for r in range(2, (ws_v.max_row or 0) + 1):
            branch = ws_v.cell(r, 1).value
            if branch not in BRANCHES:
                continue

            d = ws_v.cell(r, 2).value
            if isinstance(d, datetime):
                current = d.date()
            elif isinstance(d, date):
                current = d
            if current is None:
                continue

            day = iso(current)
            bucket = days.setdefault(day, {"pay": 0.0, "lines": [], "extras": []})

            cell_v = ws_v.cell(r, 8)
            cell_f = ws_f.cell(r, 8)
            pay = cell_v.value
            if isinstance(pay, (int, float)) and pay > 0:
                bucket["pay"] = float(pay)
                comment_items = (
                    parse_comment(cell_f.comment.text) if cell_f.comment else []
                )
                formula_nums = parse_formula_nums(cell_f.value)
                bucket["lines"] = match_lines(float(pay), formula_nums, comment_items)

            c9 = ws_v.cell(r, 9).value
            if isinstance(c9, (int, float)) and c9 > 0:
                bucket["extras"].append(float(c9))

    return days


def build_rows(days: dict[str, dict]):
    rows: list[tuple[str, str, float, str, str | None, str]] = []
    for day, info in sorted(days.items()):
        # นับชื่อซ้ำในวันเดียวกัน
        name_count: dict[str, int] = {}
        for name, amt in info["lines"]:
            name_count[name] = name_count.get(name, 0) + 1
            idx = name_count[name]
            key_suffix = slug(name) if idx == 1 else f"{slug(name)}-{idx}"
            rows.append(
                (
                    f"import:account12:note:{day}:{key_suffix}",
                    day,
                    round(amt, 2),
                    name,
                    "จากคอมเมนต์เซลล์จ่าย",
                    category_for(name),
                )
            )
        for i, amt in enumerate(info["extras"]):
            rows.append(
                (
                    f"import:account12:extra:{day}:{i + 1}",
                    day,
                    round(amt, 2),
                    "จ่ายพิเศษ",
                    None,
                    CATEGORY_SHEET,
                )
            )
    return rows


def upsert_cash_book(cur, source_key, entry_date, amount, description, note):
    cur.execute(
        "SELECT id FROM cash_book_entries WHERE source_key = %s", (source_key,)
    )
    existing = cur.fetchone()
    if existing:
        cur.execute(
            """
            UPDATE cash_book_entries SET
              entry_date=%s, direction='out', kind='expense',
              branch_id=NULL, category_id=NULL, amount=%s,
              description=%s, note=%s, updated_at=NOW(), deleted_at=NULL
            WHERE id=%s
            """,
            (entry_date, amount, description, note, existing[0]),
        )
        return
    cur.execute(
        """
        INSERT INTO cash_book_entries (
          id, entry_date, direction, kind, branch_id, category_id,
          amount, description, note, source_key, updated_at, deleted_at
        ) VALUES (%s,%s,'out','expense',NULL,NULL,%s,%s,%s,%s,NOW(),NULL)
        """,
        (
            f"cbe-{uuid.uuid4().hex[:10]}",
            entry_date,
            amount,
            description,
            note,
            source_key,
        ),
    )


def upsert_expense(cur, cash_move_id, entry_date, amount, description, note, category_id):
    spent_at = f"{entry_date}T12:00:00+07:00"
    cur.execute("SELECT id FROM expenses WHERE cash_move_id = %s", (cash_move_id,))
    existing = cur.fetchone()
    if existing:
        cur.execute(
            """
            UPDATE expenses SET
              branch_id=NULL, category_id=%s, amount=%s, description=%s,
              paid_from='external', spent_at=%s::timestamptz, note=%s,
              updated_at=NOW(), deleted_at=NULL
            WHERE id=%s
            """,
            (category_id, amount, description, spent_at, note, existing[0]),
        )
        return
    cur.execute(
        """
        INSERT INTO expenses (
          id, branch_id, category_id, amount, description, paid_from,
          spent_at, cash_move_id, note, updated_at, deleted_at
        ) VALUES (%s,NULL,%s,%s,%s,'external',%s::timestamptz,%s,%s,NOW(),NULL)
        """,
        (
            f"exp-{uuid.uuid4().hex[:10]}",
            category_id,
            amount,
            description,
            spent_at,
            cash_move_id,
            note,
        ),
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", default=str(Path.home() / "Downloads" / "Account12.xlsx"))
    ap.add_argument("--database-url", default="postgresql://appleclub@localhost:5432/ap_pos")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    days = parse_workbook(Path(args.file))
    rows = build_rows(days)

    # verify day sums vs pay
    bad = 0
    for day, info in days.items():
        s = round(sum(a for _, a in info["lines"]), 2)
        if info["pay"] and abs(s - info["pay"]) >= 1:
            bad += 1
            print("MISMATCH", day, "pay", info["pay"], "lines", s, info["lines"])

    by_desc: dict[str, float] = {}
    for r in rows:
        by_desc[r[3]] = by_desc.get(r[3], 0) + r[2]

    print(f"days={len(days)} lines={len(rows)} total={sum(r[2] for r in rows):,.2f} mismatches={bad}")
    print("top:")
    for name, amt in sorted(by_desc.items(), key=lambda x: -x[1])[:12]:
        print(f"  {name}: {amt:,.2f}")
    print("sample 2026-04-01:", days.get("2026-04-01"))

    if args.dry_run:
        return

    with psycopg.connect(args.database_url) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO expense_categories (id, name, sort_order, active, updated_at, deleted_at)
                VALUES (%s, 'จากชีทบัญชี', 90, 1, NOW(), NULL)
                ON CONFLICT (id) DO UPDATE SET active=1, deleted_at=NULL, updated_at=NOW()
                """,
                (CATEGORY_SHEET,),
            )
            keep = [r[0] for r in rows]
            for source_key, entry_date, amount, description, note, category_id in rows:
                upsert_cash_book(cur, source_key, entry_date, amount, description, note)
                upsert_expense(
                    cur, source_key, entry_date, amount, description, note, category_id
                )

            if keep:
                ph = ",".join(["%s"] * len(keep))
                cur.execute(
                    f"""
                    UPDATE cash_book_entries SET deleted_at=NOW(), updated_at=NOW()
                    WHERE deleted_at IS NULL AND source_key LIKE 'import:account12:%%'
                      AND source_key NOT IN ({ph})
                    """,
                    keep,
                )
                stale_cb = cur.rowcount
                cur.execute(
                    f"""
                    UPDATE expenses SET deleted_at=NOW(), updated_at=NOW()
                    WHERE deleted_at IS NULL AND cash_move_id LIKE 'import:account12:%%'
                      AND cash_move_id NOT IN ({ph})
                    """,
                    keep,
                )
                stale_ex = cur.rowcount
            else:
                stale_cb = stale_ex = 0

        conn.commit()

    print(f"imported={len(rows)} stale_cb={stale_cb} stale_ex={stale_ex}")


if __name__ == "__main__":
    main()
